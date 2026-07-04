"""Build a unified commissions pivot from combined_output.xlsx.

Combines individual commissions (``Final Commission Adjusted``) and managerial
commissions (``Managerial Commission``) into one pivot:

- Rows:    Person (the commission recipient)
- Columns: Sale Type -> Brand -> Product Type
- Values:  total commission (individual + managerial), summed
- Row totals per person and a grand total (``margins``)

Each sheet in combined_output.xlsx is one person's commissions. A sheet is an
individual sheet unless it carries a ``Managerial Commission`` column, in which
case the recipient is the manager (``Mapped_Name``) rather than the seller.
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INDIVIDUAL_VALUE = "Final Commission Adjusted"
MANAGERIAL_VALUE = "Managerial Commission"

CATEGORY_COLUMNS = ["Sale Type", "Brand", "Product Type"]


def to_number(series: pd.Series) -> pd.Series:
    """Parse a ``$1,234.56`` currency-string column back into floats."""
    cleaned = (
        series.astype(str)
        .str.replace(r"[$,]", "", regex=True)
        .str.strip()
        .replace({"": None, "nan": None, "None": None})
    )
    return pd.to_numeric(cleaned, errors="coerce")


def normalize_label(series: pd.Series) -> pd.Series:
    """Title-case category/person text so both sources share the same labels.

    Missing values (``NaN``/``NA``) become an empty string so they survive as a
    valid "blank" group key (e.g. rows with no Brand) instead of being dropped
    by the pivot.
    """
    return (
        series.fillna("")
        .astype(str)
        .str.strip()
        .str.title()
        .replace({"Nan": "", "None": "", "<Na>": ""})
    )


def column_or_blank(df: pd.DataFrame, column: str) -> pd.Series:
    """Return the column if present, else an all-blank Series of the right length."""
    if column in df.columns:
        return df[column]
    return pd.Series([""] * len(df), index=df.index)


def load_records(input_file: Path) -> pd.DataFrame:
    """Read every sheet and return long-format rows: Person + categories + amount."""
    records: list[pd.DataFrame] = []

    excel = pd.ExcelFile(input_file)
    for sheet in excel.sheet_names:
        df = pd.read_excel(excel, sheet_name=sheet)

        if MANAGERIAL_VALUE in df.columns:
            person_col, value_col = "Mapped_Name", MANAGERIAL_VALUE
        elif INDIVIDUAL_VALUE in df.columns:
            person_col, value_col = "Opportunity Owner", INDIVIDUAL_VALUE
        else:
            # e.g. unmatched_owners sheet - nothing to sum, skip it.
            continue

        part = pd.DataFrame(
            {
                "Person": normalize_label(column_or_blank(df, person_col)),
                "Sale Type": normalize_label(column_or_blank(df, "Sale Type")),
                "Brand": normalize_label(column_or_blank(df, "Brand")),
                "Product Type": normalize_label(column_or_blank(df, "Product Type")),
                "Commission": to_number(df[value_col]),
            }
        )
        records.append(part)

    if not records:
        raise ValueError(f"No commission sheets found in {input_file}")

    data = pd.concat(records, ignore_index=True)
    data["Commission"] = data["Commission"].fillna(0.0)
    return data


def build_pivot(data: pd.DataFrame) -> pd.DataFrame:
    """Pivot: Person x (Sale Type, Brand, Product Type), summed, with totals."""
    return pd.pivot_table(
        data,
        index="Person",
        columns=CATEGORY_COLUMNS,
        values="Commission",
        aggfunc="sum",
        fill_value=0,
        margins=True,
        margins_name="Grand Total",
    ).round(2)


def save_formatted_pivot(pivot: pd.DataFrame, output_file: Path) -> None:
    """Write the pivot to Excel with header, currency, total and freeze styling.

    Layout produced by ``to_excel`` for a 3-level column pivot: rows 1-3 are the
    Sale Type / Brand / Product Type header levels, row 4 holds the ``Person``
    index name, and data starts on row 5. The Grand Total sits in the last row
    and last column.
    """
    pivot = pivot.copy()
    pivot.index.name = "Person"
    pivot.to_excel(output_file, merge_cells=True)

    workbook = load_workbook(output_file)
    worksheet = workbook.active

    header_rows = pivot.columns.nlevels + 1  # 3 column levels + index-name row
    data_start = header_rows + 1
    last_row, last_col = worksheet.max_row, worksheet.max_column

    fill = lambda hex_color: PatternFill("solid", fgColor=hex_color)
    header_fill, sub_fill, product_fill = fill("1F3864"), fill("2E5496"), fill("8EA9DB")
    person_fill, zebra_fill, total_fill = fill("F2F2F2"), fill("F7FAFF"), fill("E8EAF6")
    white_bold = Font(color="FFFFFF", bold=True)
    navy_bold = Font(color="1F3864", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    money_format = '$#,##0.00;-$#,##0.00;""'  # blank when zero

    for row in range(1, header_rows + 1):
        for col in range(1, last_col + 1):
            cell = worksheet.cell(row, col)
            cell.border, cell.alignment = border, center
            if row == 1:
                cell.fill, cell.font = header_fill, white_bold
            elif row == 2:
                cell.fill, cell.font = sub_fill, white_bold
            elif row == 3:
                cell.fill, cell.font = product_fill, navy_bold
            else:
                cell.fill, cell.font = sub_fill, white_bold

    for row in range(data_start, last_row + 1):
        total_row = worksheet.cell(row, 1).value == "Grand Total"
        for col in range(1, last_col + 1):
            cell = worksheet.cell(row, col)
            cell.border = border
            if col == 1:
                cell.alignment = left
                cell.font = Font(bold=True, color="000000" if total_row else "1F3864")
                cell.fill = total_fill if total_row else person_fill
            else:
                cell.alignment, cell.number_format = right, money_format
                if total_row or col == last_col:
                    cell.font, cell.fill = Font(bold=True), total_fill
                elif (row - data_start) % 2 == 1:
                    cell.fill = zebra_fill

    worksheet.column_dimensions["A"].width = 24
    for col in range(2, last_col + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = 13
    for row in range(1, header_rows + 1):
        worksheet.row_dimensions[row].height = 30
    worksheet.freeze_panes = worksheet.cell(row=data_start, column=2)

    workbook.save(output_file)


def main() -> None:
    """Read combined_output.xlsx and write the unified commissions pivot."""
    entered = input(
        "Enter the combined workbook path [combined_output.xlsx]: "
    ).strip()
    input_file = Path(entered) if entered else Path("combined_output.xlsx")
    if not input_file.exists():
        raise FileNotFoundError(f"File not found: {input_file}")

    output_file = Path("commissions_pivot.xlsx")

    data = load_records(input_file)
    pivot = build_pivot(data)
    save_formatted_pivot(pivot, output_file)

    print(f"Combined {len(data)} rows across {data['Person'].nunique()} people.")
    print(f"Pivot shape (rows x cols): {pivot.shape}")
    print(f"Saved {output_file}")


if __name__ == "__main__":
    main()
